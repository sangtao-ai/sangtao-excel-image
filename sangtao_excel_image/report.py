"""Ghi ket qua ra file.

Ghi ngay sau moi job chu khong doi chay het, de mat dien giua chung van con
nguyen phan da lam. File goc cua nguoi dung khong bao gio bi sua.
"""

from __future__ import annotations

import csv
from dataclasses import asdict, dataclass
from pathlib import Path

HEADERS = [
    "dong",
    "trang_thai",
    "prompt",
    "so_anh_tham_chieu",
    "jobId",
    "link_anh",
    "file_da_tai",
    "credit",
    "loi",
]


@dataclass
class Entry:
    line: int
    status: str  # complete | error | skipped
    prompt: str
    reference_count: int = 0
    job_id: str = ""
    url: str = ""
    file: str = ""
    credit_cost: float = 0.0
    error: str = ""

    def as_row(self) -> list[object]:
        return [
            self.line,
            self.status,
            self.prompt,
            self.reference_count,
            self.job_id,
            self.url,
            self.file,
            self.credit_cost,
            self.error,
        ]


class Report:
    """Ghi dan ra CSV, ket thuc thi xuat them ban .xlsx cho de doc."""

    def __init__(self, out_dir: Path):
        self.out_dir = out_dir
        self.csv_path = out_dir / "_ket-qua.csv"
        self.entries: list[Entry] = []

        out_dir.mkdir(parents=True, exist_ok=True)
        # utf-8-sig de Excel mo len khong loi tieng Viet.
        self._fh = self.csv_path.open("w", newline="", encoding="utf-8-sig")
        self._writer = csv.writer(self._fh)
        self._writer.writerow(HEADERS)
        self._fh.flush()

    def add(self, entry: Entry) -> None:
        self.entries.append(entry)
        self._writer.writerow(entry.as_row())
        self._fh.flush()  # flush ngay, khong giu trong buffer

    def close(self) -> None:
        try:
            self._fh.close()
        except Exception:
            pass
        self._write_xlsx()

    def _write_xlsx(self) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError:
            return  # co CSV la du, khong sao

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Ket qua"

            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for entry in self.entries:
                ws.append(entry.as_row())

            widths = [6, 12, 60, 8, 34, 60, 28, 8, 50]
            for idx, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = width

            ws.freeze_panes = "A2"
            wb.save(self.out_dir / "_ket-qua.xlsx")
        except Exception:
            pass  # CSV da co roi, khong lam hong lan chay vi viec nay

    # ------------------------------------------------------------- thong ke

    @property
    def done(self) -> int:
        return sum(1 for e in self.entries if e.status == "complete")

    @property
    def failed(self) -> int:
        return sum(1 for e in self.entries if e.status == "error")

    @property
    def skipped(self) -> int:
        return sum(1 for e in self.entries if e.status == "skipped")

    @property
    def total_credit(self) -> float:
        return round(sum(e.credit_cost for e in self.entries), 2)
