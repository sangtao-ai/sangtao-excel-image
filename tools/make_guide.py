"""Sinh HUONG-DAN.txt va HUONG-DAN.docx tu cung mot noi dung.

Giu mot ban goc duy nhat de hai file khong bao gio lech nhau.
"""

from __future__ import annotations

import sys
from pathlib import Path

TITLE = "TẠO ẢNH HÀNG LOẠT TỪ FILE EXCEL — sangtao.ai"

REPO_URL = "https://github.com/sangtao-ai/sangtao-excel-image"

INTRO = (
    "Công cụ này đọc một file Excel, mỗi dòng là một ảnh cần tạo, "
    "rồi tự gọi API sangtao.ai và tải ảnh về máy bạn."
)

# Moi muc: (tieu de, [khoi noi dung])
# Khoi la chuoi thuong, hoac ("code", "..."), ("bullet", [...]), ("steps", [...])
SECTIONS: list[tuple[str, list]] = [
    ("LẦN ĐẦU SỬ DỤNG", [
        ("steps", [
            "Giải nén thư mục này ra (ví dụ ra D:\\tao-anh).\n"
            "Đừng chạy trực tiếp từ bên trong file .zip.",

            "Lấy API key: vào sangtao.ai → đăng nhập → Cài đặt → API Key → "
            "tạo key mới. Copy chuỗi key đó.",

            "Mở file mau.xlsx, điền nội dung ảnh bạn muốn tạo "
            "(xem phần CÁCH ĐIỀN FILE bên dưới), rồi lưu lại.",

            "Bấm đúp vào TaoAnh.exe.",

            "Chương trình sẽ hỏi API key. Dán key vào rồi bấm Enter.\n"
            "Key được lưu vào config.txt, lần sau không phải nhập lại.",

            "Ngồi đợi. Mỗi ảnh mất khoảng 30–90 giây.",
        ]),
        ("note", "Lần đầu chạy, Windows có thể hiện bảng cảnh báo màu xanh "
                 "\"Windows protected your PC\". Bấm \"More info\" rồi bấm "
                 "\"Run anyway\". Đây là vì file chưa được mua chứng chỉ ký số, "
                 "không phải virus. Mã nguồn công khai tại:"),
        ("code", REPO_URL),
    ]),

    ("CÁCH ĐIỀN FILE", [
        "File cần dùng 4 cột. Dòng đầu tiên là dòng tiêu đề:",
        ("code", "prompt | images | aspectRatio | format"),
        "Chỉ cột \"prompt\" là BẮT BUỘC. Ba cột còn lại để trống cũng được.",

        ("field", ("prompt", [
            "Mô tả ảnh bạn muốn tạo.",
            "Viết đủ: cảnh gì, phong cách gì, bố cục ra sao.",
            "Viết quá ngắn (ví dụ \"hình vuông xanh\") sẽ bị từ chối.",
        ])),
        ("code", "A vintage travel poster of Ha Long Bay at sunset,\n"
                 "bold retro typography, limestone karsts against an orange sky"),

        ("field", ("images", [
            "Ảnh tham chiếu — để trống nếu không cần. Tối đa 10 ảnh mỗi dòng.",
            "Điền được cả hai kiểu:",
            "    • Đường dẫn trên máy:  D:\\anh\\logo.png",
            "    • Link trên mạng:      https://example.com/anh.jpg",
            "Nhiều ảnh thì phân cách bằng dấu chấm phẩy:",
            "    D:\\anh\\logo.png; https://example.com/anh.jpg",
            "Ảnh trên máy sẽ được tự động tải lên trước khi tạo ảnh.",
            "Nếu nhiều dòng dùng chung một ảnh, nó chỉ tải lên một lần.",
        ])),

        ("field", ("aspectRatio", [
            "Tỷ lệ khung ảnh: 1:1, 16:9, 9:16...",
            "Để trống thì hệ thống tự chọn.",
        ])),

        ("field", ("format", [
            "png hoặc jpeg. Để trống thì mặc định png.",
        ])),

        ("note", "Mỗi dòng cho ra MỘT ảnh. Muốn 3 ảnh cùng nội dung thì "
                 "copy dòng đó thành 3 dòng."),
    ]),

    ("KẾT QUẢ NẰM Ở ĐÂU", [
        "Sau khi chạy xong, cạnh file Excel sẽ có một thư mục mới tên kiểu:",
        ("code", "ket-qua_2026-09-14_0930\\\n"
                 "    0001.png          ← ảnh của dòng 2 trong Excel\n"
                 "    0002.png          ← ảnh của dòng 3\n"
                 "    ...\n"
                 "    _ket-qua.xlsx     ← bảng tổng hợp"),
        "File _ket-qua.xlsx cho biết từng dòng: thành công hay lỗi, link ảnh, "
        "tên file đã tải về, và hết bao nhiêu credit.",
        ("warn", "Ảnh trên máy chủ bị xoá sau 7 ngày. Chương trình đã tải về "
                 "máy bạn rồi, nhưng đừng chỉ dựa vào link trong bảng kết quả "
                 "để lưu trữ lâu dài."),
    ]),

    ("KIỂM TRA TRƯỚC KHI TỐN TIỀN", [
        "Muốn kiểm tra file đã điền đúng chưa mà chưa muốn tạo ảnh thật:",
        "Mở Command Prompt tại thư mục này rồi gõ:",
        ("code", "TaoAnh.exe mau.xlsx --dry-run"),
        "Lệnh này kiểm tra từng dòng, báo rõ dòng nào thiếu prompt, dòng nào "
        "sai đường dẫn ảnh — và KHÔNG tạo ảnh, KHÔNG trừ tiền.",
        ("note", "Nên chạy thử cái này trước khi chạy file có vài trăm dòng."),
    ]),

    ("LỖI HAY GẶP", [
        ("error", ("API key không đúng",
                   "Key sai hoặc đã bị xoá. Mở config.txt, dán lại key mới.")),
        ("error", ("Không tìm thấy ảnh trên máy: ...",
                   "Đường dẫn trong cột images sai. Chương trình có in ra nó "
                   "đã tìm ở đâu — so lại với vị trí thật của file ảnh.\n"
                   "Mẹo: bấm chuột phải vào file ảnh → \"Copy as path\" rồi dán vào.")),
        ("error", ("PROMPT_UNCLEAR",
                   "Mô tả quá ngắn hoặc quá mơ hồ. Viết dài và cụ thể hơn.")),
        ("error", ("CONTENT_REJECTED",
                   "Nội dung vi phạm chính sách. Đổi cách mô tả.")),
        ("error", ("Không tải được ảnh tham chiếu",
                   "Link ảnh không mở được từ máy chủ. Một số trang chặn tải ảnh "
                   "từ bên ngoài — trong trường hợp đó, tải ảnh về máy rồi điền "
                   "đường dẫn trên máy thay vì điền link.")),
        ("error", ("Mở file Excel trong lúc chạy",
                   "Chương trình không đọc được. Đóng Excel lại rồi chạy lại.")),
        ("note", "Job bị lỗi KHÔNG bị trừ tiền, credit được hoàn tự động."),
    ]),

    ("CHẠY LẠI NHỮNG DÒNG BỊ LỖI", [
        "Mở _ket-qua.xlsx, xem cột \"trang_thai\" để biết dòng nào lỗi.",
        "Tạo một file Excel mới chỉ chứa những dòng đó (đã sửa lỗi), "
        "rồi chạy lại bình thường.",
        "Chương trình không bao giờ ghi đè lên thư mục kết quả cũ, nên ảnh "
        "đã tải về từ lần trước vẫn còn nguyên.",
    ]),

    ("LƯU Ý VỀ BẢO MẬT", [
        "File config.txt chứa API key của bạn. Đây là thông tin bí mật:",
        ("bullet", [
            "Đừng gửi file này cho người khác",
            "Đừng đưa lên GitHub, Google Drive dùng chung, hay nhóm chat",
            "Ai có key đều dùng được credit của bạn",
        ]),
        "Nếu lỡ lộ key: vào sangtao.ai → Cài đặt → API Key → xoá key cũ "
        "và tạo key mới.",
    ]),

    ("CÁC LỆNH KHÁC", [
        ("cmd", ("Chạy với một file cụ thể", "TaoAnh.exe D:\\duong-dan\\jobs.xlsx")),
        ("cmd", ("Chọn nơi lưu kết quả", "TaoAnh.exe jobs.xlsx -o D:\\anh-khach-A")),
        ("cmd", ("Chỉ kiểm tra, không tạo ảnh", "TaoAnh.exe jobs.xlsx --dry-run")),
        ("cmd", ("Xem chi tiết hơn khi chạy", "TaoAnh.exe jobs.xlsx -v")),
        ("cmd", ("Xem tất cả lựa chọn", "TaoAnh.exe --help")),
        ("note", "Cũng có thể kéo thả file .xlsx thẳng vào biểu tượng TaoAnh.exe."),
    ]),
]

LINE = "=" * 64
THIN = "-" * 64


def _wrap(
    text: str,
    width: int = 64,
    indent: str = "",
    hanging: str | None = None,
) -> list[str]:
    """Xuong dong theo tu, giu nguyen cac dong da thut le san.

    hanging la phan thut le cho cac dong thu hai tro di — dung cho muc danh so
    de chu thang hang duoi chu, khong thut ve sat le.
    """
    cont = indent if hanging is None else hanging
    out: list[str] = []

    for para in text.split("\n"):
        if not para.strip():
            out.append("")
            continue
        if para.startswith(("    ", "\t")) or len(indent + para) <= width:
            out.append(indent + para)
            continue

        line = indent
        first = True
        for word in para.split():
            prefix = indent if first else cont
            if line.strip() and len(line) + len(word) + 1 > width:
                out.append(line.rstrip())
                line = cont + word
                first = False
            else:
                line = f"{line} {word}" if line.strip() else prefix + word
        if line.strip():
            out.append(line.rstrip())
    return out


def build_txt(dest: Path) -> None:
    lines: list[str] = [LINE, f"  {TITLE}", LINE, ""]
    lines += _wrap(INTRO)
    lines.append("")

    for title, blocks in SECTIONS:
        lines += ["", THIN, f"  {title}", THIN, ""]

        for block in blocks:
            if isinstance(block, str):
                lines += _wrap(block)
                lines.append("")
                continue

            kind, payload = block

            if kind == "code":
                lines += [f"    {ln}" for ln in payload.split("\n")]
                lines.append("")

            elif kind == "steps":
                for index, step in enumerate(payload, start=1):
                    first, *rest = step.split("\n")
                    # hanging="   " de dong thu hai thang hang voi chu, khong
                    # tut ve duoi so thu tu.
                    lines += _wrap(f"{index}. {first}", hanging="   ")
                    for extra in rest:
                        lines += _wrap(extra, indent="   ", hanging="   ")
                    lines.append("")

            elif kind == "bullet":
                for item in payload:
                    lines += _wrap(f"  - {item}", hanging="    ")
                lines.append("")

            elif kind == "field":
                name, details = payload
                lines.append(f"  {name}")
                for detail in details:
                    lines += _wrap(detail, indent="      ", hanging="      ")
                lines.append("")

            elif kind == "error":
                name, fix = payload
                lines.append(f'  "{name}"')
                lines += _wrap(fix, indent="      ", hanging="      ")
                lines.append("")

            elif kind == "cmd":
                label, command = payload
                lines.append(f"  {label}:")
                lines.append(f"      {command}")
                lines.append("")

            elif kind in ("note", "warn"):
                prefix = "  LƯU Ý: " if kind == "note" else "  QUAN TRỌNG: "
                wrapped = _wrap(prefix + payload, hanging="  ")
                lines += wrapped
                lines.append("")

    text = "\n".join(lines).rstrip() + "\n"
    # Ghi kem BOM de Notepad tren Windows hien dung tieng Viet.
    dest.write_text(text, encoding="utf-8-sig")


def build_docx(dest: Path) -> bool:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError:
        return False

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(16)

    intro = doc.add_paragraph(INTRO)
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def code(text: str) -> None:
        for line in text.split("\n"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def callout(text: str, label: str, color: RGBColor) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(12)
        tag = para.add_run(f"{label} ")
        tag.bold = True
        tag.font.color.rgb = color
        para.add_run(text)

    for title, blocks in SECTIONS:
        doc.add_heading(title, level=1)

        for block in blocks:
            if isinstance(block, str):
                doc.add_paragraph(block)
                continue

            kind, payload = block

            if kind == "code":
                code(payload)

            elif kind == "steps":
                for step in payload:
                    doc.add_paragraph(step, style="List Number")

            elif kind == "bullet":
                for item in payload:
                    doc.add_paragraph(item, style="List Bullet")

            elif kind == "field":
                name, details = payload
                para = doc.add_paragraph()
                run = para.add_run(name)
                run.bold = True
                run.font.name = "Consolas"
                for detail in details:
                    sub = doc.add_paragraph(detail)
                    sub.paragraph_format.left_indent = Pt(24)
                    sub.paragraph_format.space_after = Pt(2)

            elif kind == "error":
                name, fix = payload
                para = doc.add_paragraph()
                run = para.add_run(f"„{name}\u201d")
                run.bold = True
                run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
                sub = doc.add_paragraph(fix)
                sub.paragraph_format.left_indent = Pt(24)

            elif kind == "cmd":
                label, command = payload
                doc.add_paragraph(label + ":")
                code(command)

            elif kind == "note":
                callout(payload, "LƯU Ý:", RGBColor(0x1F, 0x6F, 0xB2))

            elif kind == "warn":
                callout(payload, "QUAN TRỌNG:", RGBColor(0xC0, 0x39, 0x2B))

    doc.save(dest)
    return True


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent
    out.mkdir(parents=True, exist_ok=True)

    build_txt(out / "HUONG-DAN.txt")
    print(f"Da tao {out / 'HUONG-DAN.txt'}")

    if build_docx(out / "HUONG-DAN.docx"):
        print(f"Da tao {out / 'HUONG-DAN.docx'}")
    else:
        print("Bo qua .docx (chua cai python-docx: pip install python-docx)")
