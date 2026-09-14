"""Sinh file mau mau.xlsx va mau.csv."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["prompt", "images", "aspectRatio", "format"]

NOTES = {
    "prompt": (
        "BAT BUOC. Mo ta anh can tao.\n"
        "Viet du canh, phong cach, bo cuc.\n"
        "Prompt qua ngan se bi tu choi."
    ),
    "images": (
        "Tuy chon. Anh tham chieu, toi da 10.\n"
        "Dan duong dan anh tren may (D:\\anh\\logo.png)\n"
        "hoac link https://...\n"
        "Nhieu anh thi phan cach bang dau ;"
    ),
    "aspectRatio": "Tuy chon. Vi du 1:1, 16:9, 9:16.\nDe trong thi he thong tu chon.",
    "format": "Tuy chon. png hoac jpeg.\nDe trong thi mac dinh png.",
}

ROWS = [
    [
        "A vintage travel poster of Ha Long Bay at sunset, bold retro typography, "
        "limestone karsts silhouetted against an orange sky, muted 1960s color palette",
        "",
        "9:16",
        "png",
    ],
    [
        "Product photo of a matte ceramic coffee mug on a white marble countertop, "
        "soft morning daylight from the left, shallow depth of field, minimal styling",
        "",
        "1:1",
        "",
    ],
    [
        "A cozy bookshop interior in autumn, warm lamp light, wooden shelves full of "
        "books, rain on the window, soft painterly style",
        "",
        "16:9",
        "png",
    ],
    [
        "Banner for a coffee brand, roasted beans scattered on dark wood, dramatic "
        "side lighting, generous copy space on the right",
        "",
        "16:9",
        "jpeg",
    ],
]

# Cot images de trong trong file mau. Neu dien san mot duong dan vi du thi lan
# chay dau tien se bao "khong tim thay anh" va nguoi dung moi dung tuong minh
# lam sai — huong dan cach dien da nam trong ghi chu cua o tieu de.


def build_xlsx(dest: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sach anh"

    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    ws.append(HEADERS)
    for index, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        note = Comment(NOTES[name], "huong dan")
        note.width = 280
        note.height = 110
        cell.comment = note

    for row in ROWS:
        ws.append(row)

    for index in range(1, len(HEADERS) + 1):
        letter = get_column_letter(index)
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=(index <= 2))

    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 10
    ws.row_dimensions[1].height = 22
    for index in range(2, len(ROWS) + 2):
        ws.row_dimensions[index].height = 46

    ws.freeze_panes = "A2"
    wb.save(dest)


def build_csv(dest: Path) -> None:
    with dest.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(HEADERS)
        writer.writerows(ROWS)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent
    out.mkdir(parents=True, exist_ok=True)
    build_xlsx(out / "mau.xlsx")
    build_csv(out / "mau.csv")
    print(f"Da tao {out / 'mau.xlsx'}")
    print(f"Da tao {out / 'mau.csv'}")
